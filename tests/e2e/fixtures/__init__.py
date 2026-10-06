"""
E2E fixtures and sample files generator.
"""
import io
import docx

def create_sample_docx_bytes():
    """Create a structured sample resume in DOCX format."""
    doc = docx.Document()
    doc.add_heading("Zaid Alam", level=0)
    doc.add_paragraph("Email: zaid@example.com | Phone: +91 9876543210")
    doc.add_paragraph("Full Stack Software Engineer specializing in Python, Django, React, and PostgreSQL.")
    
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, Django, FastAPI, PostgreSQL, Docker, Redis, TypeScript, React, Git")
    
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Engineer at CloudTech Systems (2023 - Present)")
    doc.add_paragraph("Developed scalable REST APIs using Django REST Framework and PostgreSQL.")
    doc.add_paragraph("Deployed microservices using Docker and managed Redis caching layers.")
    
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Technology in Computer Science & Engineering")
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
